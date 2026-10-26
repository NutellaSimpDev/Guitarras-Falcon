# -*- coding: utf-8 -*-
"""Genera el Estudio de Mercado de Guitarras Falcón (docx)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\EduPedu\XiaomiMiMoProjects\Proyecto Falcon\Guitarras-Falcon\docs\Estudio-de-mercado-Guitarras-Falcon.docx"

# Paleta del proyecto (tinta café / latón)
INK = RGBColor(0x1A, 0x12, 0x0C)
ACCENT = RGBColor(0x6B, 0x3F, 0x22)
SOFT = RGBColor(0x59, 0x59, 0x59)
BRASS = RGBColor(0xC0, 0x8B, 0x3E)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.8)
    section.orientation = WD_ORIENTATION.PORTRAIT


def tune_styles(doc):
    body = doc.styles["Normal"]
    body.font.name = "Calibri"
    body.font.size = Pt(11)
    body.font.color.rgb = RGBColor(0x2A, 0x21, 0x1A)
    body.paragraph_format.line_spacing = 1.15
    body.paragraph_format.space_after = Pt(8)

    for name, size, color, before in [
        ("Title", 28, INK, 0),
        ("Subtitle", 14, SOFT, 0),
        ("Heading 1", 18, ACCENT, 16),
        ("Heading 2", 14, INK, 12),
        ("Heading 3", 12, ACCENT, 10),
    ]:
        s = doc.styles[name]
        s.font.name = "Calibri Light" if "Heading" in name or name == "Title" else "Calibri"
        s.font.size = Pt(size)
        s.font.bold = name != "Subtitle"
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(6)

    cap = doc.styles["Caption"]
    cap.font.name = "Calibri"
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.color.rgb = SOFT

    quote = doc.styles["Quote"]
    quote.font.name = "Calibri"
    quote.font.size = Pt(11)
    quote.font.italic = True
    quote.font.color.rgb = ACCENT


def add_footer(doc, text):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = SOFT
        r.font.name = "Calibri"


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph("Estudio de Mercado", style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st = doc.add_paragraph("Guitarras Falcón — Luthería peruana desde 1946", style="Subtitle")
    st.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph()
    r = p.add_run("Diagnóstico comercial, oportunidades de crecimiento y plan de acción")
    r.font.size = Pt(13)
    r.font.color.rgb = ACCENT

    doc.add_paragraph()
    for line in [
        "Elaborado para: Familia Falcón · Taller Luna Pizarro 310, La Victoria, Lima",
        "Alcance: mercado peruano de guitarras artesanales y música criolla / folclórica",
        "Documento de trabajo · Revisión con la familia antes de implementar",
    ]:
        q = doc.add_paragraph(line)
        q.paragraph_format.space_after = Pt(2)
        for r in q.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = SOFT

    doc.add_page_break()


def add_toc_placeholder(doc):
    doc.add_heading("Contenido", level=1)
    items = [
        "1. Resumen ejecutivo",
        "2. El problema que estamos resolviendo",
        "3. Contexto del mercado",
        "4. Perfil del comprador (buyer personas)",
        "5. Análisis de la competencia",
        "6. Análisis FODA de Guitarras Falcón",
        "7. Oportunidades priorizadas",
        "8. Estrategia comercial propuesta",
        "9. Plan de acción (90 días y 12 meses)",
        "10. CRM y seguimiento de ventas",
        "11. Indicadores (KPIs)",
        "12. Fuentes y notas de método",
    ]
    for i in items:
        doc.add_paragraph(i, style="List Number")
    doc.add_page_break()


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def p(doc, text):
    return doc.add_paragraph(text)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, name in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = name
        for para in cell.paragraphs:
            for r in para.runs:
                r.bold = True
    for ri, row in enumerate(rows, start=1):
        for ci, value in enumerate(row):
            t.rows[ri].cells[ci].text = str(value)
    doc.add_paragraph()
    return t


def build():
    doc = Document()
    setup_page(doc)
    tune_styles(doc)
    add_footer(doc, "Guitarras Falcón · Estudio de mercado · Documento de trabajo interno")
    add_cover(doc)
    add_toc_placeholder(doc)

    # ---------- 1 ----------
    h1(doc, "1. Resumen ejecutivo")
    p(doc, (
        "Guitarras Falcón es una marca con un activo difícil de replicar: casi ocho décadas "
        "de historia, una estructura interna patentada (INDECOPI, Resolución N.º 014132-2009) "
        "y el respaldo emocional de los grandes nombres de la música criolla y andina del Perú. "
        "El problema actual no es de producto, sino de visibilidad y de sistema comercial: "
        "las ventas han bajado porque la tradición musical se ha enfriado entre los más jóvenes, "
        "la marca no está presente donde hoy se decide la compra (Google, WhatsApp, redes sociales) "
        "y no existe un proceso ordenado para captar y dar seguimiento a los interesados."
    ))
    p(doc, (
        "Este estudio propone tres movimientos coordinados: (1) una web moderna que cuente la "
        "historia y convierta visitas en conversaciones de WhatsApp; (2) una estrategia de "
        "contenido y alianzas que reactive el interés por la guitarra peruana; y (3) un CRM "
        "sencillo para no perder ningún lead. El objetivo a 12 meses es aumentar las consultas "
        "cualificadas y la conversión a venta, sin depender de publicidad cara."
    ))

    h2(doc, "Conclusiones clave")
    bullets(doc, [
        "El producto es defendible: patente, sonido reconocido y una narrativa única en el país.",
        "El cuello de botella está en la captación y el seguimiento, no en la fabricación.",
        "WhatsApp es el canal de cierre natural en Perú; la web debe alimentarlo, no competir con él.",
        "La música criolla tiene reconocimiento oficial (Patrimonio Cultural de la Nación, 2022) pero atraviesa una «resaca» generacional: hay que ir a buscar al público joven en sus canales.",
        "Las vías de crecimiento más realistas a corto plazo son: turismo cultural, educación musical, regalo corporativo / institucional y comunidad de peñas.",
    ])

    # ---------- 2 ----------
    h1(doc, "2. El problema que estamos resolviendo")
    p(doc, (
        "La familia ha notado una caída de ventas en los últimos años. El diagnóstico preliminar, "
        "a partir de la información del taller y del comportamiento del mercado, apunta a tres causas combinadas:"
    ))
    bullets(doc, [
        "Cambio generacional: los jóvenes consumen reguetón, trap, pop y K-pop; la guitarra criolla dejó de ser el «instrumento de entrada» natural.",
        "Canales: la mayoría de luthiers y talleres peruanos siguen vendiendo de boca en boca o con webs informativas sin embudo de venta.",
        "Sin sistema de seguimiento: los interesados llegan por WhatsApp, redes o visita al taller, pero no queda registro estructurado para darles seguimiento.",
    ])
    p(doc, (
        "La buena noticia es que el activo de marca (Abraham Falcón, los músicos que tocaron estas "
        "guitarras, la patente) es exactamente el tipo de historia que hoy funciona bien en internet "
        "cuando se cuenta bien. El problema no es que la tradición no valga: es que casi nadie la ha "
        "contado con las herramientas actuales."
    ))

    # ---------- 3 ----------
    h1(doc, "3. Contexto del mercado")
    h2(doc, "3.1. La música criolla en el Perú")
    p(doc, (
        "La música criolla y afroperuana es un pilar de la identidad costeña del Perú. El vals "
        "peruano, la marinera, la polca criolla, el festejo o el tondero forman un ecosistema en el "
        "que la guitarra es el instrumento protagonista. En 2022 la música criolla fue reconocida "
        "como Patrimonio Cultural de la Nación (Lima y Callao), y cada 31 de octubre se celebra el "
        "Día de la Canción Criolla."
    ))
    p(doc, (
        "Al mismo tiempo, especialistas como el musicólogo Marino Martínez Espinoza han señalado "
        "que el género atravesó una «resaca» por la desaparición de cultores y estilos de canto. "
        "En 2024, el Archivo Nacional de Derecho de Autor (INDECOPI) custodiaba más de 70.000 obras "
        "del criollismo, incluidas casi 10.000 grabaciones sonoras: hay patrimonio, pero falta "
        "relevo vivo. Para Guitarras Falcón esto significa que el mercado no desapareció: se está "
        "reconfigurando y necesita nuevos intermediarios (docentes, peñas, creadores de contenido, "
        "turismo cultural)."
    ))

    h2(doc, "3.2. El mercado de guitarras en el Perú")
    p(doc, (
        "No existe una estadística oficial y actualizada del tamaño del mercado de guitarras "
        "artesanales en el Perú. A efectos de planificación, se propone segmentarlo así:"
    ))
    table(doc,
        ["Segmento", "Qué busca", "Ticket aprox. (S/)", "Papel de Falcón"],
        [
            ["Estudiante / iniciación", "Precio accesible, comodidad, resistencia", "600 – 1.500", "Línea Estudio"],
            ["Aficionado / peña", "Volumen, carácter criollo, uso real", "1.200 – 2.500", "Export Peña (foco)"],
            ["Músico avanzado / docente", "Proyección, equilibrio, prestigio", "2.500 – 4.500+", "Semi-Profesional"],
            ["Coleccionista / regalo", "Historia, patente, valor simbólico", "2.000 – 5.000+", "Ediciones y piezas con relato"],
            ["Institucional", "Precio por volumen, plazos, factura", "Variable", "Colegios, universidades, peñas, institutos"],
            ["Turismo / exportación", "Relato cultural, portabilidad, caja de regalo", "2.000 – 6.000+", "Piezas de souvenir premium"],
        ]
    )
    p(doc, (
        "Nota de método: los tickets son estimaciones de trabajo basadas en el rango típico de "
        "guitarras criollas y de luthier en el mercado limeño. Conviene validarlos con el histórico "
        "real del taller (precios de las últimas 20 ventas) antes de fijar la lista de precios."
    ))

    h2(doc, "3.3. Tendencias a favor")
    bullets(doc, [
        "El «efecto documental»: las historias de oficio y patrimonio tienen buena recepción en YouTube, Instagram y TikTok.",
        "La venta consultiva por WhatsApp es el estándar en el Perú; no hace falta carrito de compra para vender bien.",
        "El turismo experiencial en Lima (Barranco, Miraflores, La Victoria, peñas) busca productos con historia real.",
        "Las escuelas de música, academias de canto criollo y peñas siguen comprando instrumentos, aunque en volúmenes menores que en los años 80.",
    ])

    # ---------- 4 ----------
    h1(doc, "4. Perfil del comprador (buyer personas)")
    p(doc, (
        "Estos perfiles están construidos a partir de los tipos de guitarra del catálogo, el "
        "formulario de contacto actual (que ya pregunta por modelo y «cómo nos conociste») y la "
        "naturaleza del taller. Sirven para orientar mensajes y prioridades de contacto."
    ))

    h2(doc, "4.1. «Doña Carmen», 55 años — mamá o abuela que regala")
    bullets(doc, [
        "Motivación: regalar una guitarra «de verdad» a un hijo o nieto que empieza a tocar.",
        "Criterios: precio razonable, confianza en el vendedor, que no se rompa.",
        "Canal: recomendación, WhatsApp, visita al taller con un familiar.",
        "Mensaje ganador: «la primera guitarra de tu hijo, hecha a mano en Lima desde 1946».",
    ])

    h2(doc, "4.2. «Luis», 28 años — músico de peña o agrupación")
    bullets(doc, [
        "Motivación: guitarra que se escuche fuerte en la jarana y aguante el uso.",
        "Criterios: sonido, resistencia, que la reconozcan los colegas.",
        "Canal: Instagram, grupos de WhatsApp de peñas, recomendación de otros músicos.",
        "Mensaje ganador: «el sonido que tocó el Zambo Cavero, listo para tu próxima peña».",
    ])

    h2(doc, "4.3. «Rosa», 38 años — docente de música")
    bullets(doc, [
        "Motivación: comprar varias guitarras para sus alumnos o su instituto.",
        "Criterios: precio por volumen, consistencia, factura, plazos.",
        "Canal: correo institucional, LinkedIn, ferias educativas.",
        "Mensaje ganador: «lote para tu academia, con el respaldo de la guitarra peruana patentada».",
    ])

    h2(doc, "4.4. «Martin», 45 años — turista o comprador internacional")
    bullets(doc, [
        "Motivación: llevarse un objeto de cultura peruana auténtico, no un souvenir de aeropuerto.",
        "Criterios: historia, autenticidad, empaque, facilidad de transporte.",
        "Canal: hoteles, agencias de turismo cultural, Instagram en inglés/español, web.",
        "Mensaje ganador: «la guitarra peruana con patente estatal, hecha en el taller donde nació el sonido criollo».",
    ])

    # ---------- 5 ----------
    h1(doc, "5. Análisis de la competencia")
    p(doc, (
        "El mercado peruano de guitarras se puede agrupar en cuatro bloques. Guitarras Falcón no "
        "compite solo con otros luthiers: compite con el tiempo de ocio del joven, con la guitarra "
        "importada barata y con la percepción de que «la criolla ya no es moderna»."
    ))
    table(doc,
        ["Bloque", "Ejemplos de oferta", "Fortaleza", "Debilidad frente a Falcón"],
        [
            ["Luthiers de autor", "Talleres artesanales en Lima y provincias", "Oficio real, relación cercana", "Poca historia institucional y patente"],
            ["Marcas importadas de entrada", "Guitarras de fábrica asiáticas", "Precio bajo, disponibilidad", "Sin identidad criolla ni relato"],
            ["Tiendas de instrumentos", "Retail multi-marca", "Variedad, financiamiento", "Venden producto, no cultura"],
            ["Mercado de segunda mano", "Facebook Marketplace, WhatsApp", "Precio muy bajo", "Sin garantía ni servicio"],
        ]
    )
    h2(doc, "Ventaja diferencial de Falcón")
    bullets(doc, [
        "Estructura patentada ante INDECOPI (única en su categoría en el Perú).",
        "Archivo histórico con los grandes de la música criolla y andina.",
        "Reconocimiento internacional (Francia, años 80, «modelo peruano»).",
        "Taller físico con historia visitable en La Victoria.",
        "Tres líneas de producto que cubren desde el estudiante hasta el concierto.",
    ])

    # ---------- 6 ----------
    h1(doc, "6. Análisis FODA de Guitarras Falcón")
    table(doc,
        ["", "Positivo", "Negativo"],
        [
            ["Interno",
             "· Casi 80 años de trayectoria\n· Estructura patentada\n· Tercera generación al frente\n· Materiales y oficio propios\n· Reconocimiento de músicos históricos",
             "· Capacidad de producción limitada (hecho a mano)\n· Presencia digital recién renovada\n· Sin CRM ni proceso comercial formal\n· Dependencia del boca a boca"],
            ["Externo",
             "· Música criolla como Patrimonio Cultural (2022)\n· Demanda de turismo experiencial\n· WhatsApp como canal de cierre masivo\n· Interés por lo artesanal y lo local\n· Alianzas con peñas, escuelas y festivales",
             "· Caída del interés juvenil por el género\n· Competencia de instrumentos importados baratos\n· Poder adquisitivo ajustado\n· Coste de las maderas\n· Escasa cobertura mediática del oficio"],
        ]
    )

    # ---------- 7 ----------
    h1(doc, "7. Oportunidades priorizadas")
    p(doc, "Se ordenan por impacto esperado y facilidad de ejecución con los recursos actuales del taller:")

    h2(doc, "7.1. Corto plazo (0–3 meses) — alto impacto, bajo coste")
    bullets(doc, [
        "Publicar la nueva web y hacerla el embudo central: toda conversación termina en WhatsApp.",
        "Activar el CRM local para registrar cada consulta y dar seguimiento en 24–48 h.",
        "Serie de contenidos «La historia detrás del sonido»: 8–12 piezas cortas (video/foto) con el archivo familiar.",
        "Catálogo PDF actualizado de los tres modelos, con fotos reales y «consultar precio».",
        "Plantillas de respuesta de WhatsApp para cada tipo de lead (estudio, peña, institucional, turista).",
    ])

    h2(doc, "7.2. Medio plazo (3–6 meses) — crecimiento sostenido")
    bullets(doc, [
        "Alianzas con 5–10 peñas y escuelas de música criolla: descuento por convenio + visita al taller.",
        "Visitas guiadas al taller (30–45 min) para turistas y colegios; posible entrada con compra de accesorio o pieza pequeña.",
        "Presencia en 2–3 ferias culturales o educativas al año (Día de la Canción Criolla, ferias de luthería, festivales de marinera).",
        "Programa «Guitarra Falcón en la escuela»: condiciones especiales para docentes y academias.",
        "Refresco de redes: Instagram y TikTok con el oficio, no solo el producto terminado.",
    ])

    h2(doc, "7.3. Largo plazo (6–12 meses) — consolidación de marca")
    bullets(doc, [
        "Edición de coleccionista «Abraham Falcón» con certificado de autenticidad y número de serie.",
        "Exportación de 1–2 piezas al mes a compradores en EE. UU., Europa y Chile (comunidad peruana).",
        "Archivo digital público (fotos, audios, documentos) que posiciona a la marca como institución cultural.",
        "Curso o taller de iniciación a la guitarra criolla en el local (genera comunidad y captación de leads).",
    ])

    # ---------- 8 ----------
    h1(doc, "8. Estrategia comercial propuesta")

    h2(doc, "8.1. Producto")
    bullets(doc, [
        "Mantener las tres líneas (Estudio, Export Peña, Semi-Profesional) como columnas del catálogo.",
        "Añadir una cuarta oferta simbólica: «Edición Archivo» o piezas con historia documentada (no implica nuevo modelo, implica nuevo relato).",
        "Incluir siempre: certificado de autenticidad, número de serie y una tarjeta con la historia de la estructura patentada.",
        "Pack turismo / regalo: guitarra + funda + cuerdas de repuesto + folleto de la historia.",
    ])

    h2(doc, "8.2. Precios")
    bullets(doc, [
        "Publicar el rango de cada línea en la web («desde S/ x» o «consultar») para filtrar leads no cualificados.",
        "Mantener el cierre de precio en WhatsApp o en el taller: la consultoría es parte del valor.",
        "Descuento de convenio (8–12 %) para peñas, academias y docentes, a cambio de referencia y visibilidad.",
        "Sin descuentos agresivos de temporada: protege el valor de la marca.",
    ])

    h2(doc, "8.3. Canales")
    table(doc,
        ["Canal", "Rol", "Acción concreta"],
        [
            ["Web (falcontradiciones.com / GitHub Pages)", "Captación y relato", "SEO local, WhatsApp integrado, formulario → CRM"],
            ["WhatsApp", "Cierre y seguimiento", "Plantillas, respuesta < 24 h, catálogo en PDF"],
            ["Instagram / TikTok", "Notoriedad y comunidad", "2–3 piezas de contenido semanales (oficio, archivo)"],
            ["Visita al taller", "Confianza y venta consultiva", "Agenda por WhatsApp, experiencia de 30–45 min"],
            ["Alianzas (peñas, escuelas, hoteles)", "Demanda recurrente", "Convenios simples con comisión o descuento"],
            ["Ferias y festivales", "Presencia física de marca", "2–3 eventos/año con demo de sonido"],
        ]
    )

    h2(doc, "8.4. Mensajes de marca")
    bullets(doc, [
        "Eje 1 — Legado: «Desde 1946, la guitarra que sonó en las grandes voces del Perú».",
        "Eje 2 — Innovación: «La única estructura de guitarra peruana patentada en el país».",
        "Eje 3 — Oficio: «Hecha a mano en La Victoria, por la tercera generación de la familia Falcón».",
        "Eje 4 — Comunidad: «No es solo una guitarra: es tu lugar en la tradición criolla».",
    ])

    # ---------- 9 ----------
    h1(doc, "9. Plan de acción")
    h2(doc, "9.1. Primeros 90 días")
    table(doc,
        ["Semana", "Acción", "Responsable", "Resultado esperado"],
        [
            ["1–2", "Publicar la nueva web y activar el CRM", "Familia / apoyo digital", "Embudo digital operativo"],
            ["1–2", "Fijar precios de referencia de las 3 líneas", "Taller", "Lista de precios interna"],
            ["2–4", "Serie de 8 contenidos con archivo familiar", "Familia + apoyo", "Banco de contenido inicial"],
            ["3–5", "Plantillas de WhatsApp y protocolo de respuesta < 24 h", "Taller", "Menos leads perdidos"],
            ["4–8", "Contacto con 10 peñas / academias para convenios", "Comercial", "5 conversaciones activas"],
            ["8–12", "Revisión de leads del CRM y ajuste de mensajes", "Familia", "Mensajes afinados por perfil"],
        ]
    )

    h2(doc, "9.2. De 3 a 12 meses")
    bullets(doc, [
        "Mensual: revisar el CRM (leads nuevos, tasa de conversión, origen) y publicar al menos 4 contenidos.",
        "Trimestral: 1 actividad externa (feria, visita guiada o alianza firmada).",
        "Semestral: ajustar precios y catálogo según demanda real observada en el CRM.",
        "Anual: evaluar edición de coleccionista y primeras ventas de exportación.",
    ])

    # ---------- 10 ----------
    h1(doc, "10. CRM y seguimiento de ventas")
    p(doc, (
        "Se ha desarrollado un prototipo de CRM ligero (incluido en este mismo proyecto, bajo /crm) "
        "pensado para el día a día del taller, sin licencias ni servidores:"
    ))
    bullets(doc, [
        "Captura automática de los formularios de la web (nombre, teléfono, modelo, origen del lead).",
        "Pipeline de estados: Nuevo → Contactado → Cotizado → Negociación → Cerrado.",
        "Vista kanban y vista tabla, búsqueda por nombre, modelo u origen.",
        "Exportación a CSV para compartir con un contador, un asesor o para abrirlo en Excel.",
        "Datos guardados en el navegador del equipo del taller (localStorage). Es una herramienta de arranque: si el volumen crece, se migra a un CRM en la nube.",
    ])
    p(doc, (
        "Regla de oro: ningún lead sin registrar. Cada conversación de WhatsApp que parezca "
        "compra potencial debe entrar al CRM ese mismo día, con el modelo de interés y la "
        "próxima acción. La conversión de un negocio de luthier se pierde en el seguimiento "
        "olvidado, no en el precio."
    ))

    # ---------- 11 ----------
    h1(doc, "11. Indicadores (KPIs)")
    table(doc,
        ["Indicador", "Definición", "Meta sugerida a 12 meses"],
        [
            ["Leads nuevos / mes", "Contactos que entran al CRM", "≥ 12"],
            ["Tiempo de primera respuesta", "Horas hasta el primer mensaje", "< 24 h"],
            ["Tasa de conversión", "Cerrados / total de leads", "15–25 %"],
            ["Origen dominante", "Qué canal trae más ventas", "WhatsApp + web ≥ 50 %"],
            ["Ticket promedio", "S/ promedio por venta", "Según mix de líneas"],
            ["Ventas institucionales", "Academias / peñas / colegios", "≥ 4 cuentas activas"],
            ["Contenido publicado", "Piezas / mes en redes", "≥ 4"],
            ["Visitantes de la web", "Sesiones / mes", "Crecimiento trimestral sostenido"],
        ]
    )

    # ---------- 12 ----------
    h1(doc, "12. Fuentes y notas de método")
    p(doc, (
        "Este documento combina (a) información del propio negocio extraída de la web y la "
        "historia familiar de Guitarras Falcón; (b) contexto público sobre la música criolla "
        "peruana; y (c) estimaciones de trabajo para segmentos y tickets, que deben validarse "
        "con el histórico real de ventas del taller."
    ))
    bullets(doc, [
        "Guitarras Falcón — sitio oficial: https://falcontradiciones.com/ (historia, productos, contacto).",
        "INDECOPI — Resolución N.º 014132-2009: registro de la «Innovación Estructural para la Guitarra Peruana Clásica de Concierto» (20 de agosto de 2009).",
        "Música criolla y afroperuana — Wikipedia en español (contexto histórico; reconocimiento como Patrimonio Cultural de la Nación en 2022; archivo de más de 70.000 obras en INDECOPI al 2024).",
        "Taller Luna Pizarro 310, La Victoria, Lima — datos de contacto: contacto@falcontradiciones.com · +51 997 842 063.",
    ])
    p(doc, (
        "Recomendación final: antes de fijar precios públicos, cruzar este estudio con el libro "
        "de ventas o las facturas de los últimos 24 meses. Si no existe registro, empezar a "
        "anotar en el CRM a partir de hoy: la primera función del CRM no es vender, es aprender."
    ))

    doc.save(OUT)
    print("OK", OUT)


if __name__ == "__main__":
    build()
